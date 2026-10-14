"""Editable PowerPoint generation (python-pptx). All numbers come from the saved result files."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1E, 0x29, 0x3B)
BLUE = RGBColor(0x25, 0x63, 0xEB)
ORANGE = RGBColor(0xD9, 0x77, 0x06)
GREY = RGBColor(0x4B, 0x55, 0x63)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xDC, 0x26, 0x26)

SW, SH = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.n = 0

    # ------------------------------------------------------------ primitives
    def slide(self, title: str, subtitle: str | None = None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        self.text(s, title, 0.45, 0.12, 12.4, 0.7, size=26, bold=True, color=WHITE)
        if subtitle:
            self.text(s, subtitle, 0.45, 1.0, 12.4, 0.45, size=14, color=GREY, italic=True)
        self.text(
            s,
            f"UCS321 · MST Mini Project #6 · PMSM winding-temperature virtual sensor · {self.n}",
            0.45,
            7.08,
            12.4,
            0.3,
            size=9,
            color=GREY,
        )
        return s

    def text(self, s, txt, x, y, w, h, size=16, bold=False, color=GREY, italic=False, align=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = txt
        run.font.size, run.font.bold, run.font.italic = Pt(size), bold, italic
        run.font.color.rgb = color
        if align:
            p.alignment = align
        return tb

    def bullets(self, s, items: list[str], x, y, w, h, size=16, color=GREY):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            lvl = 1 if it.startswith("  ") else 0
            p.level = lvl
            r = p.add_run()
            r.text = ("• " if lvl == 0 else "– ") + it.strip()
            r.font.size = Pt(size - 2 * lvl)
            r.font.color.rgb = color
            p.space_after = Pt(5)
        return tb

    def image(self, s, path: Path, x, y, w=None, h=None):
        if not Path(path).is_file():
            self.text(s, f"[missing figure: {Path(path).name}]", x, y, 4, 0.4, color=RED)
            return None
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)

    def callout(self, s, txt, x, y, w, h, fill=LIGHT, color=NAVY, size=14, bold=False):
        box = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.color.rgb = BLUE
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = txt
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
        return box

    def table(self, s, df: pd.DataFrame, x, y, w, col_w: list[float] | None = None, size=11):
        rows, cols = df.shape[0] + 1, df.shape[1]
        shape = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(0.34 * rows))
        tbl = shape.table
        for j, c in enumerate(df.columns):
            cell = tbl.cell(0, j)
            cell.text = str(c) or " "
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size, para.runs[0].font.bold = Pt(size), True
            para.runs[0].font.color.rgb = WHITE
        for i in range(df.shape[0]):
            for j in range(cols):
                cell = tbl.cell(i + 1, j)
                cell.text = str(df.iloc[i, j]) or " "
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
                para = cell.text_frame.paragraphs[0]
                para.runs[0].font.size = Pt(size)
                para.runs[0].font.color.rgb = NAVY
        if col_w:
            for j, cw in enumerate(col_w):
                tbl.columns[j].width = Emu(int(Inches(cw)))
        return tbl

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


def f(x: float, nd: int = 2) -> str:
    return f"{x:.{nd}f}"


def build_deck(ctx: dict[str, Any], out: Path) -> Path:
    """Compose the 12-slide deck from a context dict assembled in :mod:`report`."""
    d = Deck()
    fig = ctx["fig_dir"]
    eda = ctx["eda_dir"]
    r, A, B = ctx["res"], ctx["res"]["track_A"], ctx["res"]["track_B"]
    am, bm = A["test"]["aggregate"], B["test"]["aggregate"]
    mode_note = "" if ctx["mode"] == "full" else "  (FAST-MODE preliminary results)"

    # 1 title
    s = d.prs.slides.add_slide(d.prs.slide_layouts[6])
    d.n += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    d.text(
        s,
        "Electric Motor Winding Temperature Estimation and Thermal Monitoring "
        "Using Machine Learning",
        0.8,
        1.6,
        11.7,
        2.0,
        size=38,
        bold=True,
        color=WHITE,
    )
    d.text(
        s,
        "A history-aware virtual sensor for a permanent-magnet synchronous motor",
        0.8,
        3.7,
        11.7,
        0.6,
        size=20,
        color=RGBColor(0xBF, 0xDB, 0xFE),
    )
    d.text(
        s,
        "UCS321 · MST Mini Project #6 · General Electric (GE) Power Systems case study",
        0.8,
        4.6,
        11.7,
        0.5,
        size=16,
        color=WHITE,
    )
    d.text(
        s,
        "[Your name] · [Roll number] · [Group / Instructor]",
        0.8,
        5.2,
        11.7,
        0.5,
        size=16,
        color=RGBColor(0xFD, 0xE6, 0x8A),
    )
    d.text(
        s,
        "Dataset: Paderborn University LEA test bench (public, CC BY-SA 4.0) - not GE hardware.",
        0.8,
        6.6,
        11.7,
        0.4,
        size=12,
        color=RGBColor(0xBF, 0xDB, 0xFE),
        italic=True,
    )

    # 2 problem
    s = d.slide(
        "Problem and motivation", "Why estimate a temperature we could in principle measure?"
    )
    d.bullets(
        s,
        [
            "Winding insulation ageing and demagnetisation limit motor torque; the controller must know "
            "how hot the stator winding is.",
            "Direct sensing: accurate, but extra cost, wiring, failure points and hard to retrofit.",
            "Virtual sensing: reuse signals the drive already has (speed, torque, d-q currents/voltages, "
            "coolant, ambient) and learn the mapping from data.",
            "Research question: can ML estimate winding temperature on unseen operating profiles "
            "without the winding temperature as input?",
            "Also: use the estimate as a reference - a persistent residual can flag something unusual "
            "(experimental, not a protection function).",
        ],
        0.5,
        1.5,
        7.2,
        5.2,
        size=17,
    )
    d.callout(
        s,
        "Direct sensor\n+ measured truth\n- cost / failure point\n\nVirtual sensor\n+ no hardware\n"
        "+ redundancy / plausibility check\n- statistical, limited to training domain",
        8.1,
        1.7,
        4.7,
        3.6,
        size=16,
    )
    d.callout(
        s,
        "GE Power Systems is the case-study framing of the assignment. The data comes from a "
        "university test bench, not GE equipment.",
        8.1,
        5.5,
        4.7,
        1.2,
        fill=RGBColor(0xFE, 0xF3, 0xC7),
        size=13,
    )

    # 3 data
    p = ctx["prov"]
    s = d.slide(
        "Dataset and provenance",
        "Paderborn University LEA PMSM test bench - Kaggle 'electric-motor-temperature'",
    )
    d.bullets(
        s,
        [
            f"measures_v2.csv · {p['file_size_bytes'] / 1e6:.0f} MB · {p['rows']:,} rows · {p['profiles']} operating "
            f"profiles · ≈ {p['total_hours_at_assumed_rate']:.0f} h at an assumed 2 Hz",
            "Target: stator_winding. Inputs: ambient, coolant, motor_speed, torque, i_d, i_q, u_d, u_q.",
            "Never used as inputs: stator_tooth, stator_yoke, pm (internal temperatures), profile_id.",
            "No missing values, no duplicate rows; each profile is one contiguous block (row order = time).",
            f"{ctx['n_warm']} of {p['profiles']} profiles start with a warm motor - not all recordings are cold starts.",
            f"SHA-256: {p['sha256'][:16]}… · licence CC BY-SA 4.0 (attribution in README).",
        ],
        0.5,
        1.5,
        6.6,
        5.4,
        size=15,
    )
    d.image(s, eda / "profile_durations.png", 7.2, 1.6, w=5.8)
    d.image(s, eda / "start_temperature.png", 8.4, 4.3, w=3.8)

    # 4 EDA
    s = d.slide("Exploratory data analysis", "What the measurements look like")
    d.image(s, eda / "example_timeseries.png", 0.4, 1.4, w=6.6)
    d.image(s, eda / "correlation_heatmap.png", 7.2, 1.4, h=4.3)
    d.bullets(
        s,
        [
            "Winding temperature spans ≈ 19-141 °C; instantaneous |i|² explains only part of it.",
            "Internal temperatures (tooth, yoke, magnet) are near-collinear with the target → excluded as inputs.",
            "Speed-torque coverage differs between profiles → split by whole profile.",
        ],
        0.5,
        5.5,
        12.4,
        1.6,
        size=14,
    )

    # 5 architecture + validation
    m = ctx["manifest"]
    s = d.slide("Pipeline and leakage-safe validation")
    d.image(s, ctx["fig_root"] / "architecture.png", 0.3, 1.1, w=8.2)
    d.bullets(
        s,
        [
            f"Split by whole profile: {m['train']['n_profiles']} train / {m['val']['n_profiles']} val / "
            f"{m['test']['n_profiles']} test profiles ({m['train']['row_fraction'] * 100:.0f} / "
            f"{m['val']['row_fraction'] * 100:.0f} / {m['test']['row_fraction'] * 100:.0f} % of rows).",
            "Stratified by peak winding temperature; manifest saved.",
            "Scaler and statistics fitted on training rows only (Pipeline).",
            "Features are causal and reset at every profile start.",
            "Test profiles opened only after the model selection is frozen.",
            "Unit tests fail if any of this breaks.",
        ],
        8.6,
        1.3,
        4.5,
        5.6,
        size=14,
    )
    d.callout(
        s,
        "Why not a random row split? Samples are 0.5 s apart and nearly identical to their neighbours, so a "
        "random split would put a near-copy of every test row into training and reward memorisation. "
        "Whole operating profiles are held out instead.",
        0.4,
        5.2,
        8.0,
        1.5,
        size=14,
    )

    # 6 features
    s = d.slide("Feature engineering and thermal inertia", "Physics → features")
    d.callout(
        s,
        "C·dT/dt = P − (T − T_coolant)/R   ⇒   T[k+1] = (1−α)·T[k] + α·(T_coolant + R·P[k]),  α = Δt/τ\n"
        "The temperature is an exponentially weighted average (EWMA) of the heating history.",
        0.5,
        1.5,
        12.3,
        1.1,
        size=15,
        bold=True,
    )
    d.bullets(
        s,
        [
            "Track A (instantaneous): 8 raw signals + 9 derived: i², |i|, |u|, |speed|, |torque|, "
            "P_mech (exact), P_elec / P_loss (proxies), T_amb − T_cool.",
            "Track B (history-aware): + causal EWMAs (τ = 30 s, 2 min, 10 min, 30 min, 2 h) and rolling means "
            "(60 s, 5 min) of i², |speed|, |torque| and coolant temperature → 45 features.",
            "Windows defined in seconds, computed on the full 2 Hz series, per profile, no future samples.",
            "Cold start is explicit: rolling means are NaN until the window is full (LightGBM handles NaN); "
            "no arbitrary fill values.",
            "Offline (vectorised) and streaming implementations are unit-tested to agree.",
        ],
        0.5,
        2.8,
        12.3,
        4.0,
        size=16,
    )

    # 7 models + results table
    cand = ctx["cand_test"]
    s = d.slide("Models compared" + mode_note, "Held-out test profiles · MAE / RMSE in °C")
    keep = [
        "A/mean_baseline/basic",
        "A/linear/physical",
        "A/ridge/enhanced",
        "A/random_forest/enhanced",
        "A/hist_gb/enhanced",
        "A/lightgbm/enhanced",
        r["selection"]["winners"]["A"],
        "B/lightgbm/enhanced+ewma_long",
        "B/hist_gb/enhanced+ewma_long_rolling",
        r["selection"]["winners"]["B"],
    ]
    tab = cand[cand.candidate.isin(keep)].drop_duplicates("candidate")
    tab = (
        tab.set_index("candidate")
        .loc[[k for k in dict.fromkeys(keep) if k in set(tab.candidate)]]
        .reset_index()
    )
    show = pd.DataFrame(
        {
            "Candidate": tab["candidate"],
            "Track": tab["track"],
            "MAE": tab["test_mae"].map(f),
            "RMSE": tab["test_rmse"].map(f),
            "R²": tab["test_r2"].map(lambda v: f(v, 3)),
            "p95 AE": tab["test_p95_ae"].map(f),
            "Worst-profile MAE": tab["test_worst_profile_mae"].map(f),
            "Fit [s]": tab["fit_seconds"].map(lambda v: f(v, 1)),
        }
    )
    d.table(s, show, 0.4, 1.5, 12.5, [4.1, 0.7, 0.9, 0.9, 0.9, 1.0, 2.1, 1.0], size=12)
    d.bullets(
        s,
        [
            "Instantaneous models plateau far above history-aware ones - algorithm choice matters much less "
            "than representing thermal memory.",
            "Winners were chosen on validation profiles only; the test table is a single pass.",
        ],
        0.5,
        5.6,
        12.3,
        1.4,
        size=14,
    )

    # 8 result plots
    s = d.slide(
        "Results on unseen profiles" + mode_note, "Track B history-aware vs Track A instantaneous"
    )
    d.image(s, fig / "scatter_track_b.png", 0.3, 1.25, h=3.3)
    d.image(s, fig / "scatter_track_a.png", 4.5, 1.25, h=3.3)
    d.image(s, fig / "per_profile_mae.png", 8.5, 1.3, w=4.6)
    kv = pd.DataFrame(
        {
            "Metric": ["MAE [°C]", "RMSE [°C]", "R²", "median AE", "p95 AE", "worst profile MAE"],
            "Track A": [
                f(am["mae"]),
                f(am["rmse"]),
                f(am["r2"], 3),
                f(am["median_ae"]),
                f(am["p95_ae"]),
                f(A["test"]["profile_level"]["worst_profile_mae"]),
            ],
            "Track B": [
                f(bm["mae"]),
                f(bm["rmse"]),
                f(bm["r2"], 3),
                f(bm["median_ae"]),
                f(bm["p95_ae"]),
                f(B["test"]["profile_level"]["worst_profile_mae"]),
            ],
        }
    )
    d.table(s, kv, 0.5, 4.75, 6.2, [2.4, 1.9, 1.9], size=11)
    d.bullets(
        s,
        [
            f"Test: {len(r['test_profiles'])} whole profiles, {r['test_rows']:,} samples.",
            "R² alone hides °C-level errors - read MAE, p95 and the worst profile.",
        ],
        7.0,
        4.9,
        6.0,
        1.8,
        size=14,
    )

    # 9 ablation + explainability
    s = d.slide(
        "Does history help? What does the model use?",
        "Ablation (validation) and permutation importance",
    )
    d.image(s, fig / "history_ablation.png", 0.3, 1.4, w=6.4)
    d.image(s, fig / "importance_groups_track_b.png", 6.8, 1.5, w=6.2)
    d.bullets(
        s,
        [
            "Same train/validation profiles for every row of the ablation; mean ± std over seeds.",
            "Importance by feature family (correlated features are shuffled together); "
            "importance is predictive, not causal.",
            ctx["ablation_sentence"],
        ],
        0.5,
        5.2,
        12.3,
        1.8,
        size=14,
    )

    # 10 monitoring
    mt = r["monitoring"]
    s = d.slide("Thermal residual monitoring (experimental)", "residual = measured − estimated")
    d.image(s, fig / "monitoring_typical_profile.png", 0.3, 1.4, w=6.3)
    d.bullets(
        s,
        [
            f"Band ±{mt['test']['band_smoothed_c']:.1f} °C: {mt['band_quantile'] * 100:.1f}th percentile of the "
            f"60 s-smoothed |residual| on validation profiles; alarm needs ≥ {mt['persistence_s']} s persistence.",
            f"On healthy test profiles: {mt['test']['persistent_alarm_frac'] * 100:.2f} % of time in alarm, "
            f"{mt['test']['profiles_with_any_alarm']} of {mt['test']['n_profiles']} profiles raise ≥ 1 alarm "
            "(= false alarms / model error).",
            "Synthetic +offset sensitivity: " + ctx["inj_sentence"],
        ],
        6.8,
        1.5,
        6.3,
        3.6,
        size=14,
    )
    d.callout(
        s,
        "Not fault detection: no labelled faults exist, so no detection accuracy is claimed. "
        "A persistent residual may mean a sensor issue, cooling change, unseen operating point, "
        "model error or a real defect. Not a certified protection system.",
        0.5,
        5.4,
        12.3,
        1.3,
        fill=RGBColor(0xFE, 0xF3, 0xC7),
        size=14,
    )

    # 11 app + extension
    s = d.slide(
        "Application and extension", "Streamlit dashboard · 60 s-ahead forecasting experiment"
    )
    shot = ctx["fig_root"] / "app_screenshot.png"
    d.image(s, shot, 0.3, 1.3, w=7.3)
    d.bullets(
        s,
        [
            "Sections: overview, dataset exploration, instantaneous prediction, profile replay, model "
            "evaluation, thermal monitoring.",
            "Loads the saved models + metadata; never fakes predictions; warns about out-of-range inputs.",
            "Replay is historical data, clearly labelled.",
            ctx["forecast_sentence"],
            "Run: uv run streamlit run app/streamlit_app.py",
        ],
        7.8,
        1.4,
        5.3,
        5.4,
        size=14,
    )

    # 12 conclusion
    s = d.slide("Conclusions, limitations and references")
    d.bullets(
        s,
        [
            ctx["conclusion_sentence"],
            "Limitations: one test bench / motor; cold-start and warm-restart effects; coolant behaviour is "
            "operator-controlled on the bench; trees do not extrapolate; no fault labels.",
            "Future work: sequence models for comparison, warm-start initialisation, per-motor calibration, "
            "uncertainty estimates, validation on real fault data.",
            "References: Kirchgässner, Wallscheid, Böcker - Paderborn LEA PMSM temperature data & deep-pmsm; "
            "Ke et al., LightGBM (NIPS 2017); scikit-learn model-validation and common-pitfalls guides.",
        ],
        0.5,
        1.3,
        12.3,
        5.6,
        size=17,
    )
    return d.save(out)
